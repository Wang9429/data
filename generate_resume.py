from pathlib import Path
import re
from xml.sax.saxutils import escape

import pymupdf
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
)


ROOT = Path(__file__).parent
SOURCE = ROOT / "王宜静-风险管理简历-JD优化版.md"
DOCX_OUT = ROOT / "王宜静-风险管理简历-JD优化版.docx"
PDF_OUT = ROOT / "王宜静-风险管理简历-JD优化版.pdf"
FONT_PATH = "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"
FONT_NAME = "WenQuanYiMicroHei"
BLUE = "1F4E79"
LIGHT_BLUE = "DCE6F1"
GRAY = "555555"


def blocks():
    result = []
    kpmg_count = 0
    for raw in SOURCE.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("### 毕马威企业咨询"):
            kpmg_count += 1
            if kpmg_count == 2:
                result.append(("pagebreak", ""))
        if line.startswith("# "):
            result.append(("name", line[2:]))
        elif line.startswith("## "):
            result.append(("section", line[3:]))
        elif line.startswith("### "):
            result.append(("company", line[4:]))
        elif line.startswith("#### "):
            result.append(("subsection", line[5:]))
        elif line.startswith("- "):
            result.append(("bullet", line[2:]))
        elif line.startswith("**") and line.endswith("**"):
            result.append(("target", line[2:-2]))
        elif line.startswith("**") and "　" in line:
            result.append(("role", line))
        else:
            result.append(("body", line))
    return result


def add_inline_runs(paragraph, text):
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        run = paragraph.add_run(part[2:-2] if bold else part)
        run.bold = bold
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def set_cell_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(OxmlElement("w:keepNext"))


def build_docx(items):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.15)
    section.bottom_margin = Cm(1.05)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.5)

    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(9.2)
    normal.paragraph_format.space_after = Pt(2.2)
    normal.paragraph_format.line_spacing = 1.08

    for kind, text in items:
        if kind == "pagebreak":
            doc.add_page_break()
            continue
        if kind == "name":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(19)
            run.font.color.rgb = RGBColor.from_string(BLUE)
        elif kind == "target":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(10.5)
        elif kind == "section":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(1.5)
            set_cell_shading(p, LIGHT_BLUE)
            set_keep_with_next(p)
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor.from_string(BLUE)
        elif kind == "company":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2.5)
            p.paragraph_format.space_after = Pt(0)
            set_keep_with_next(p)
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(9.3)
            run.font.color.rgb = RGBColor.from_string(BLUE)
        elif kind == "role":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            set_keep_with_next(p)
            add_inline_runs(p, text)
            for run in p.runs:
                run.font.size = Pt(9.1)
                run.font.color.rgb = RGBColor.from_string(GRAY)
        elif kind == "subsection":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(0.5)
            set_keep_with_next(p)
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(9.2)
            run.font.color.rgb = RGBColor.from_string(BLUE)
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.35)
            p.paragraph_format.first_line_indent = Cm(-0.2)
            add_inline_runs(p, text)
        else:
            p = doc.add_paragraph()
            if "18482179729" in text:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(3)
                for run_text in text.split("｜"):
                    run = p.add_run(run_text.strip())
                    if run_text != text.split("｜")[-1]:
                        p.add_run("  |  ")
            else:
                add_inline_runs(p, text)

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("王宜静｜风险管理岗位简历")
        run.font.size = Pt(7)
        run.font.color.rgb = RGBColor(120, 120, 120)
    doc.save(DOCX_OUT)


def rich_text(text):
    escaped = escape(text)
    return re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", escaped)


def build_pdf(items):
    pdfmetrics.registerFont(TTFont(FONT_NAME, FONT_PATH))
    styles = getSampleStyleSheet()
    common = dict(fontName=FONT_NAME, textColor=colors.HexColor("#222222"))
    style = {
        "name": ParagraphStyle(
            "Name", parent=styles["Normal"], fontSize=18, leading=20,
            alignment=TA_CENTER, textColor=colors.HexColor(f"#{BLUE}"),
            spaceAfter=1.5 * mm, **{"fontName": FONT_NAME}
        ),
        "target": ParagraphStyle(
            "Target", parent=styles["Normal"], fontSize=10.5, leading=13,
            alignment=TA_CENTER, spaceAfter=0.8 * mm, **common
        ),
        "contact": ParagraphStyle(
            "Contact", parent=styles["Normal"], fontSize=8.8, leading=11,
            alignment=TA_CENTER, spaceAfter=2 * mm, **common
        ),
        "section": ParagraphStyle(
            "Section", parent=styles["Normal"], fontSize=11, leading=13.5,
            textColor=colors.HexColor(f"#{BLUE}"), backColor=colors.HexColor(f"#{LIGHT_BLUE}"),
            spaceBefore=1.4 * mm, spaceAfter=0.8 * mm, keepWithNext=True,
            borderPadding=(1.2 * mm, 1.2 * mm, 0.8 * mm), **{"fontName": FONT_NAME}
        ),
        "company": ParagraphStyle(
            "Company", parent=styles["Normal"], fontSize=10, leading=12,
            textColor=colors.HexColor(f"#{BLUE}"), spaceBefore=1.3 * mm,
            spaceAfter=0.3 * mm, keepWithNext=True, **{"fontName": FONT_NAME}
        ),
        "role": ParagraphStyle(
            "Role", parent=styles["Normal"], fontSize=8.9, leading=11,
            textColor=colors.HexColor(f"#{GRAY}"), spaceAfter=0.7 * mm,
            keepWithNext=True, **{"fontName": FONT_NAME}
        ),
        "subsection": ParagraphStyle(
            "Subsection", parent=styles["Normal"], fontSize=9.3, leading=11.4,
            textColor=colors.HexColor(f"#{BLUE}"), spaceBefore=0.8 * mm,
            spaceAfter=0.3 * mm, keepWithNext=True, **{"fontName": FONT_NAME}
        ),
        "body": ParagraphStyle(
            "Body", parent=styles["Normal"], fontSize=9.1, leading=12.7,
            spaceAfter=0.9 * mm, alignment=0, **common
        ),
        "bullet": ParagraphStyle(
            "Bullet", parent=styles["Normal"], fontSize=8.95, leading=12.6,
            leftIndent=3.4 * mm, firstLineIndent=-2.5 * mm, bulletIndent=0.7 * mm,
            spaceAfter=0.75 * mm, **common
        ),
    }

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(FONT_NAME, 6.8)
        canvas.setFillColor(colors.HexColor("#777777"))
        canvas.drawCentredString(A4[0] / 2, 7 * mm, f"王宜静｜风险管理岗位简历｜{doc.page}")
        canvas.restoreState()

    doc = BaseDocTemplate(
        str(PDF_OUT), pagesize=A4, leftMargin=12.5 * mm, rightMargin=12.5 * mm,
        topMargin=10 * mm, bottomMargin=11 * mm, title="王宜静-风险管理简历-JD优化版",
        author="王宜静",
    )
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    doc.addPageTemplates(PageTemplate(id="resume", frames=[frame], onPage=footer))

    story = []
    for kind, text in items:
        if kind == "pagebreak":
            story.append(PageBreak())
        elif kind == "name":
            story.append(Paragraph(f"<b>{escape(text)}</b>", style["name"]))
        elif kind == "target":
            story.append(Paragraph(f"<b>{escape(text)}</b>", style["target"]))
        elif kind == "section":
            story.append(Paragraph(f"<b>{escape(text)}</b>", style["section"]))
        elif kind == "company":
            story.append(Paragraph(f"<b>{escape(text)}</b>", style["company"]))
        elif kind == "role":
            story.append(Paragraph(rich_text(text), style["role"]))
        elif kind == "subsection":
            story.append(Paragraph(f"<b>{escape(text)}</b>", style["subsection"]))
        elif kind == "bullet":
            story.append(Paragraph(rich_text(text), style["bullet"], bulletText="•"))
        else:
            use_style = style["contact"] if "18482179729" in text else style["body"]
            story.append(Paragraph(rich_text(text), use_style))
    doc.build(story)


if __name__ == "__main__":
    content = blocks()
    build_docx(content)
    build_pdf(content)
    pdf = pymupdf.open(PDF_OUT)
    print(f"Generated {DOCX_OUT.name}")
    print(f"Generated {PDF_OUT.name}: {pdf.page_count} pages")
