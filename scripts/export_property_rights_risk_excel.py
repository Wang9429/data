#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Export property rights penetration risk library (26 scenarios) to Excel."""

import csv
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).parent))
from property_rights_risk_scenarios import L1_NAMES, REQ_SOURCES, SCENARIOS

SRC_CSV = Path("/home/ubuntu/.cursor/projects/workspace/uploads/_____71bd.csv")
OUT_XLSX = Path("/workspace/cnooc-risk-platform/data/产权管理领域穿透式监管风险库.xlsx")
OUT_CSV = Path("/workspace/cnooc-risk-platform/data/property-rights-penetration-risk-library.csv")
ROOT_XLSX = Path("/workspace/产权管理领域穿透式监管风险库.xlsx")
UPLOAD_XLSX = Path("/home/ubuntu/.cursor/projects/workspace/uploads/产权管理领域穿透式监管风险库.xlsx")

COLUMNS = [
    "序号",
    "风险场景\n（以国资委46号文为主）",
    "子风险场景",
    "风险点描述",
    "监管要求/风险事件",
    "穿透监管必要性",
    "目前是否有管控举措/管控是否已到位\n（待调研具体了解）",
    "穿透层级",
    "穿透层级描述",
    "主要管控阶段",
    "基于制度的要求（现状）",
    "流程控制要求/管控规则",
    "KRI及监控规则",
    "系统实现逻辑",
]


def load_original_data():
    """Load risk descriptions and system requirements from original library."""
    desc_map = {}
    req_map = {}
    with SRC_CSV.open("r", encoding="utf-8") as f:
        for row in csv.reader(f):
            if len(row) >= 11 and row[2].strip() and row[2] != "子风险场景":
                name = row[2].strip()
                desc_map[name] = row[3].strip() if len(row) > 3 else ""
                req_map[name] = row[10].strip() if len(row) > 10 else ""
    return desc_map, req_map


def build_system_req(scenario_id, req_map):
    parts = []
    for src_name in REQ_SOURCES.get(scenario_id, []):
        text = req_map.get(src_name, "")
        if text and text not in parts:
            parts.append(text)
    return "\n\n".join(parts)


def build_risk_desc(scenario_id, desc_map, fallback=""):
    parts = []
    for src_name in REQ_SOURCES.get(scenario_id, []):
        text = desc_map.get(src_name, "")
        if text and text not in parts:
            parts.append(text)
    if parts:
        return "\n".join(parts)
    return fallback


def build_rows():
    desc_map, req_map = load_original_data()
    rows = []
    for i, s in enumerate(SCENARIOS, 1):
        l1 = L1_NAMES[s["l1_key"]]
        rows.append({
            "序号": i,
            "风险场景\n（以国资委46号文为主）": l1,
            "子风险场景": s["l2"],
            "风险点描述": build_risk_desc(s["id"], desc_map, s.get("risk_desc", "")),
            "监管要求/风险事件": "",
            "穿透监管必要性": s["necessity"],
            "目前是否有管控举措/管控是否已到位\n（待调研具体了解）": "",
            "穿透层级": s["level"],
            "穿透层级描述": s["level_desc"],
            "主要管控阶段": s["stage"],
            "基于制度的要求（现状）": build_system_req(s["id"], req_map),
            "流程控制要求/管控规则": s["process"],
            "KRI及监控规则": s["kri"],
            "系统实现逻辑": s["system"],
        })
    return rows


def style_excel(writer, df):
  sheet_name = "产权管理穿透式监管风险库"
  df.to_excel(writer, index=False, sheet_name=sheet_name)
  ws = writer.sheets[sheet_name]
  header_font = Font(bold=True)
  wrap = Alignment(wrap_text=True, vertical="top")
  for col_idx, col_name in enumerate(df.columns, 1):
      cell = ws.cell(row=1, column=col_idx)
      cell.font = header_font
      cell.alignment = wrap
      width = 12
      if "风险场景" in col_name:
          width = 28
      elif col_name in ("子风险场景", "穿透层级"):
          width = 18
      elif col_name in ("风险点描述", "穿透监管必要性", "穿透层级描述"):
          width = 30
      elif col_name in ("主要管控阶段", "流程控制要求/管控规则", "系统实现逻辑"):
          width = 35
      elif col_name in ("基于制度的要求（现状）", "KRI及监控规则"):
          width = 40
      ws.column_dimensions[get_column_letter(col_idx)].width = width
  for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
      for cell in row:
          cell.alignment = wrap
  ws.freeze_panes = "A2"


def main():
    rows = build_rows()
    df = pd.DataFrame(rows, columns=COLUMNS)
    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
        style_excel(writer, df)
    df.to_csv(OUT_CSV, index=False, encoding="utf-8-sig")
    import shutil
    shutil.copy2(OUT_XLSX, ROOT_XLSX)
    UPLOAD_XLSX.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUT_XLSX, UPLOAD_XLSX)
    print(f"Exported {len(rows)} scenarios")
    print(f"  Excel: {OUT_XLSX}")
    print(f"  CSV:   {OUT_CSV}")
    print(f"  Copy:  {ROOT_XLSX}")


if __name__ == "__main__":
    main()
